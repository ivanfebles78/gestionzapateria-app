import io
import re
import zipfile
from datetime import date as date_cls
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.api.dependencies import require_admin
from app.core.config import settings
from app.db.deps import get_db
from app.models import DailyAttachment, DailyExpense, DailySale, MonthlyExpense, User

router = APIRouter(prefix='/api/exports', tags=['exports'])


KIND_LABELS = {
    'ticket_manana': 'TicketManana',
    'ticket_cierre': 'TicketCierre',
    'gasto': 'Gasto',
    'otro': 'Otro',
}


def _slug(value: str) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9_-]+', '', (value or '').replace(' ', ''))
    return cleaned or 'archivo'


def _resolve_date_range(
    month_key: str | None,
    date_from: date_cls | None,
    date_to: date_cls | None,
) -> tuple[date_cls, date_cls]:
    if month_key:
        if not re.fullmatch(r'\d{4}-\d{2}', month_key):
            raise HTTPException(status_code=400, detail='month_key inválido (YYYY-MM)')
        year, month = int(month_key[:4]), int(month_key[5:7])
        start = date_cls(year, month, 1)
        if month == 12:
            end = date_cls(year + 1, 1, 1)
        else:
            end = date_cls(year, month + 1, 1)
        # end exclusivo → último día inclusive
        last_day = date_cls.fromordinal(end.toordinal() - 1)
        return start, last_day
    if date_from and date_to:
        if date_from > date_to:
            raise HTTPException(status_code=400, detail='date_from posterior a date_to')
        return date_from, date_to
    raise HTTPException(status_code=400, detail='Indica month_key o date_from y date_to')


@router.get('/sales.xlsx')
def export_sales_xlsx(
    month_key: str | None = Query(default=None),
    date_from: date_cls | None = Query(default=None),
    date_to: date_cls | None = Query(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise HTTPException(status_code=500, detail='openpyxl no está instalado en el backend') from exc

    start, end = _resolve_date_range(month_key, date_from, date_to)

    sales = (
        db.query(DailySale)
        .filter(DailySale.sale_date >= start, DailySale.sale_date <= end)
        .order_by(DailySale.sale_date.asc())
        .all()
    )
    expenses_by_date: dict[date_cls, list[DailyExpense]] = {}
    for exp in (
        db.query(DailyExpense)
        .filter(DailyExpense.sale_date >= start, DailyExpense.sale_date <= end)
        .order_by(DailyExpense.sale_date.asc(), DailyExpense.created_at.asc())
        .all()
    ):
        expenses_by_date.setdefault(exp.sale_date, []).append(exp)

    monthly_rows = []
    if month_key:
        monthly_rows = (
            db.query(MonthlyExpense)
            .filter(MonthlyExpense.month_key == month_key)
            .order_by(MonthlyExpense.category.asc())
            .all()
        )

    wb = Workbook()
    ws_sales = wb.active
    ws_sales.title = 'Ventas'

    headers = [
        'Fecha', 'Día',
        'Mañana Efectivo', 'Mañana Tarjeta', 'Mañana Bizum', 'Mañana Bonos', 'Mañana Total',
        'Mañana Clientes',
        'Tarde Efectivo', 'Tarde Tarjeta', 'Tarde Bizum', 'Tarde Bonos', 'Tarde Total',
        'Tarde Clientes',
        'Total día', 'Gastos día', 'Balance día', 'Clientes día',
    ]
    ws_sales.append(headers)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='0F172A')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for cell in ws_sales[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    weekday_es = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
    for s in sales:
        ws_sales.append([
            s.sale_date.isoformat(),
            weekday_es[s.sale_date.weekday()],
            float(s.morning_cash), float(s.morning_card), float(s.morning_bizum), float(s.morning_bonos),
            float(s.morning_total),
            int(s.morning_customers_total),
            float(s.afternoon_cash), float(s.afternoon_card), float(s.afternoon_bizum), float(s.afternoon_bonos),
            float(s.afternoon_total),
            int(s.afternoon_customers_total),
            float(s.total_sales), float(s.daily_expenses_total), float(s.daily_balance),
            int(s.customers_total),
        ])

    for col_idx, _ in enumerate(headers, start=1):
        ws_sales.column_dimensions[ws_sales.cell(row=1, column=col_idx).column_letter].width = 16

    ws_exp = wb.create_sheet('Gastos diarios')
    ws_exp.append(['Fecha', 'Concepto', 'Importe'])
    for cell in ws_exp[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    for sale_date, items in sorted(expenses_by_date.items()):
        for exp in items:
            ws_exp.append([sale_date.isoformat(), exp.concept, float(exp.amount)])
    ws_exp.column_dimensions['A'].width = 14
    ws_exp.column_dimensions['B'].width = 40
    ws_exp.column_dimensions['C'].width = 14

    if monthly_rows:
        ws_month = wb.create_sheet('Gastos mensuales')
        ws_month.append(['Mes', 'Categoría', 'Importe'])
        for cell in ws_month[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        for row in monthly_rows:
            ws_month.append([row.month_key, row.category, float(row.amount)])
        ws_month.column_dimensions['A'].width = 12
        ws_month.column_dimensions['B'].width = 32
        ws_month.column_dimensions['C'].width = 14

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename_range = month_key or f"{start.isoformat()}_{end.isoformat()}"
    filename = f"ventas_{filename_range}.xlsx"
    headers_resp = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return StreamingResponse(
        buffer,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers=headers_resp,
    )


@router.get('/attachments.zip')
def export_attachments_zip(
    month_key: str | None = Query(default=None),
    date_from: date_cls | None = Query(default=None),
    date_to: date_cls | None = Query(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    start, end = _resolve_date_range(month_key, date_from, date_to)

    records = (
        db.query(DailyAttachment)
        .filter(DailyAttachment.sale_date >= start, DailyAttachment.sale_date <= end)
        .order_by(DailyAttachment.sale_date.asc(), DailyAttachment.created_at.asc())
        .all()
    )

    base_dir = Path(settings.UPLOAD_DIR)
    buffer = io.BytesIO()
    used_names: dict[str, int] = {}

    with zipfile.ZipFile(buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        for rec in records:
            file_path = base_dir / rec.stored_filename
            if not file_path.is_file():
                continue
            ext = Path(rec.stored_filename).suffix or Path(rec.original_filename or '').suffix or '.bin'
            kind_label = KIND_LABELS.get(rec.kind, _slug(rec.kind))
            d = rec.sale_date
            base_name = f"{d.day:02d}_{d.month:02d}_{d.year:04d}_{kind_label}"
            count = used_names.get(base_name, 0)
            used_names[base_name] = count + 1
            suffix = '' if count == 0 else f"_{count + 1}"
            archive_name = f"{base_name}{suffix}{ext.lower()}"
            zf.write(file_path, arcname=archive_name)

    buffer.seek(0)
    filename_range = month_key or f"{start.isoformat()}_{end.isoformat()}"
    filename = f"adjuntos_{filename_range}.zip"
    return StreamingResponse(
        buffer,
        media_type='application/zip',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )
